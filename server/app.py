import base64,hashlib,hmac,io,os,secrets
from datetime import date,datetime,timedelta,timezone
from functools import wraps
from pathlib import Path
from urllib.parse import quote
import bcrypt
import psycopg
from psycopg.types.json import Jsonb
from flask import Flask,request,jsonify,g,send_file,send_from_directory
from openpyxl import Workbook,load_workbook
from db import connection
from domain import *

app=Flask(__name__,static_folder=None)
app.json.ensure_ascii=False
ORIGIN=os.environ.get('APP_ORIGIN','http://localhost:8080').rstrip('/')
SECURE=ORIGIN.startswith('https://')
ROOT=Path(os.environ.get('STATIC_ROOT',str(Path(__file__).resolve().parent.parent/'dist')))
ARCHIVED="(r.archived_at IS NOT NULL OR (r.decision_date IS NOT NULL AND r.registered_at + interval '2 months' <= (now() AT TIME ZONE 'Europe/Moscow')))"
STAGE="CASE WHEN r.decision_date IS NOT NULL THEN 'done' WHEN r.uoso_executor_id IS NOT NULL THEN 'uoso_work' WHEN r.fd_status IS NOT NULL THEN 'uoso_wait' WHEN r.fd_executor_id IS NOT NULL THEN 'fd_work' ELSE 'fd_wait' END"

def digest(s):return hashlib.sha256(s.encode()).hexdigest()
def js(v):return Jsonb(v,dumps=lambda x:__import__('json').dumps(x,ensure_ascii=False,default=str))
def payload():
    value=request.get_json(silent=True)
    if not isinstance(value,dict):raise Invalid('Ожидается JSON-объект')
    return value
def roles(*permitted):
    def decorator(fn):
        @wraps(fn)
        def wrapped(*a,**kw):
            if g.user['role'] not in permitted:raise PermissionError('Недостаточно прав')
            return fn(*a,**kw)
        return wrapped
    return decorator

@app.before_request
def security():
    if not request.path.startswith('/api/') or request.path=='/api/health':return
    if request.method not in ('GET','HEAD','OPTIONS'):
        if request.headers.get('Origin')!=ORIGIN:
            return jsonify(error='Недопустимый источник запроса'),403
    if request.path=='/api/auth/login':return
    token=request.cookies.get('skz_session','')
    if not token:return jsonify(error='Необходим вход в систему'),401
    with connection() as db:
        u=db.execute('SELECT u.*,s.csrf_token_hash FROM sessions s JOIN users u ON u.id=s.user_id WHERE s.token_hash=%s AND s.expires_at>now() AND u.active',(digest(token),)).fetchone()
    if not u:return jsonify(error='Сессия завершена. Войдите снова.'),401
    g.user=u
    if request.method not in ('GET','HEAD','OPTIONS') and not secrets.compare_digest(digest(request.headers.get('X-CSRF-Token','')),u['csrf_token_hash']):
        return jsonify(error='Неверный CSRF-токен. Обновите страницу.'),403
    if u['must_change_password'] and request.path not in ('/api/me','/api/auth/change-password','/api/auth/logout'):
        return jsonify(error='Необходимо сменить пароль',mustChange=True),403

@app.after_request
def headers(response):
    response.headers['X-Content-Type-Options']='nosniff'
    response.headers['X-Frame-Options']='DENY'
    response.headers['Referrer-Policy']='same-origin'
    response.headers['Cache-Control']='no-store' if request.path.startswith('/api/') or request.path in ('/','/config.js') else 'no-cache'
    response.headers['Content-Security-Policy']="default-src 'self'; script-src 'self'; style-src 'self' 'unsafe-inline' https://fonts.googleapis.com; font-src 'self' https://fonts.gstatic.com; img-src 'self' data:; connect-src 'self'; frame-ancestors 'none'; base-uri 'self'; form-action 'self'"
    if SECURE:response.headers['Strict-Transport-Security']='max-age=31536000'
    return response

@app.errorhandler(Invalid)
def invalid(e):return jsonify(error=str(e)),400
@app.errorhandler(PermissionError)
def forbidden(e):return jsonify(error=str(e)),403
@app.errorhandler(psycopg.errors.UniqueViolation)
def duplicate(e):return jsonify(error='Значение уже существует. Обновите страницу.'),409
@app.errorhandler(psycopg.errors.CheckViolation)
def constraint(e):return jsonify(error='Данные не соответствуют требованиям. Проверьте обязательные поля.'),400
@app.errorhandler(Exception)
def unexpected(e):
    from werkzeug.exceptions import HTTPException
    if isinstance(e,HTTPException):return jsonify(error=e.description),e.code
    app.logger.exception('Request failed')
    return jsonify(error='Не удалось выполнить операцию. Повторите попытку.'),500

@app.get('/api/health')
def health():
    with connection() as db:db.execute('SELECT 1')
    return jsonify(status='ok')

def new_session(db,uid):
    token=secrets.token_urlsafe(32);csrf=csrf_for(token)
    db.execute('INSERT INTO sessions(token_hash,user_id,csrf_token_hash,expires_at) VALUES (%s,%s,%s,now()+interval \'12 hours\')',(digest(token),uid,digest(csrf)))
    return token,csrf
def csrf_for(token):
    secret=os.environ.get('SESSION_SECRET','')
    if len(secret)<32 or secret.lower().startswith(('replace','change')):raise RuntimeError('SESSION_SECRET must contain at least 32 non-placeholder characters')
    return hmac.new(secret.encode(),token.encode(),hashlib.sha256).hexdigest()
def set_session(response,token):
    response.set_cookie('skz_session',token,httponly=True,secure=SECURE,samesite='Strict',max_age=43200,path='/')
    return response
def hash_password(password):
    if not isinstance(password,str) or len(password)<8 or len(password.encode())>72:raise Invalid('Пароль: минимум 8 символов, максимум 72 байта UTF-8')
    return bcrypt.hashpw(password.encode(),bcrypt.gensalt(rounds=12)).decode()

@app.post('/api/auth/login')
def login():
    data=payload();name=str(data.get('login',''));password=data.get('password','')
    if not isinstance(password,str) or len(password.encode())>72:return jsonify(error='Неверный логин или пароль'),401
    with connection() as db:
        u=db.execute('SELECT * FROM users WHERE login=%s',(name,)).fetchone()
        valid=u and u['active'] and bcrypt.checkpw(password.encode(),u['password_hash'].encode())
        if not valid:return jsonify(error='Неверный логин или пароль'),401
        token,csrf=new_session(db,u['id'])
    return set_session(jsonify(user=public_user(u),csrf=csrf),token)

@app.get('/api/me')
def me():
    return jsonify(user=public_user(g.user),csrf=csrf_for(request.cookies['skz_session']))

@app.post('/api/auth/change-password')
def change_password():
    data=payload();old=data.get('current','');new=data.get('password','')
    if not isinstance(old,str) or len(old.encode())>72 or not bcrypt.checkpw(old.encode(),g.user['password_hash'].encode()):raise Invalid('Текущий пароль неверен')
    if old==new:raise Invalid('Новый пароль должен отличаться от прежнего')
    hashed=hash_password(new)
    with connection() as db:
        db.execute('UPDATE users SET password_hash=%s,must_change_password=false WHERE id=%s',(hashed,g.user['id']))
        db.execute('DELETE FROM sessions WHERE user_id=%s',(g.user['id'],))
        token,csrf=new_session(db,g.user['id'])
        audit(db,'Смена пароля',after={'userId':g.user['id']})
    return set_session(jsonify(csrf=csrf,ok=True),token)

@app.post('/api/auth/logout')
def logout():
    with connection() as db:db.execute('DELETE FROM sessions WHERE token_hash=%s',(digest(request.cookies.get('skz_session','')),))
    response=jsonify(ok=True);response.delete_cookie('skz_session',path='/');return response

def visibility(user,alias='r'):
    role=user['role'];rules=[];args=[]
    if role!=0:rules.append('NOT '+ARCHIVED)
    if role==1:rules.append('r.uploaded_by=%s');args.append(user['id'])
    if role==4:rules.append("r.fd_executor_id=%s AND r.fd_status IS DISTINCT FROM 'Проведена'");args.append(user['id'])
    if role==5:rules.append('r.uoso_executor_id=%s AND r.decision_date IS NULL');args.append(user['id'])
    return (' AND '.join(rules) or 'TRUE'),args

def query_filters(user,q,include_archive=False):
    where,args=visibility(user)
    if q.get('archive')=='1':
        if user['role']!=0:raise PermissionError('Архив доступен только Администратору')
        where+=' AND '+ARCHIVED
    elif not include_archive:where+=' AND NOT '+ARCHIVED
    for key,col in [('query','request_id'),('citizen','citizen_id')]:
        if q.get(key):where+=' AND position(%s in r.'+col+')>0';args.append(q[key])
    if q.get('from'):where+=' AND r.registered_at>=%s';args.append(parse_date(q['from']))
    if q.get('to'):where+=" AND r.registered_at<%s::date+interval '1 day'";args.append(parse_date(q['to']))
    if q.get('tab')=='completed':where+=' AND r.decision_date IS NOT NULL'
    if q.get('tab')=='unfinished':where+=' AND r.decision_date IS NULL'
    if q.get('stage'):
        if user['role']==1:raise PermissionError('Этапы недоступны')
        where+=' AND ('+STAGE+')=%s';args.append(q['stage'])
    return where,args

def get_record(db,reqid,lock=False):
    r=db.execute('SELECT r.* FROM requests r WHERE r.request_id=%s'+(' FOR UPDATE' if lock else ''),(reqid,)).fetchone()
    if not r or not can_view(r,g.user):raise PermissionError('Заявка недоступна')
    return r
def holidays(db):return {r['date'] for r in db.execute('SELECT date FROM holidays')}
def audit(db,action,rid=None,before=None,after=None):
    db.execute('INSERT INTO action_log(request_id,user_id,action,before_values,after_values) VALUES (%s,%s,%s,%s,%s)',(rid,g.user['id'],action,js(before),js(after)))
def notify(db,rid,role,message,recipient=None):
    if not recipient:
        u=db.execute('SELECT id FROM users WHERE role=%s AND active ORDER BY id LIMIT 1',(role,)).fetchone();recipient=u['id'] if u else None
    if recipient:db.execute('INSERT INTO notifications(recipient_id,request_id,message) VALUES (%s,%s,%s)',(recipient,rid,message))

@app.get('/api/state')
def state_data():
    with connection() as db:
        if g.user['role']==0:us=db.execute('SELECT * FROM users ORDER BY id').fetchall()
        elif g.user['role'] in (2,3):us=db.execute('SELECT * FROM users WHERE id=%s OR role IN (4,5) ORDER BY id',(g.user['id'],)).fetchall()
        else:us=[g.user]
        hs=[r['date'].isoformat() for r in db.execute('SELECT date FROM holidays')]
    return jsonify(user=public_user(g.user),users=[public_user(u) for u in us],holidays=hs)

@app.get('/api/requests')
def list_requests():
    where,args=query_filters(g.user,request.args)
    sort={'2':'request_id','3':'registered_at','12':'fd_status','22':'status'}.get(request.args.get('sort'),'registered_at')
    if g.user['role']==1 and sort in ('fd_status','status'):sort='registered_at'
    direction='DESC' if request.args.get('desc','1')=='1' else 'ASC'
    try:page=max(1,int(request.args.get('page',1)))
    except ValueError:raise Invalid('Некорректная страница')
    with connection() as db:
        recipient='TRUE' if g.user['role']==0 else 'n.recipient_id='+str(int(g.user['id']))
        new=f"EXISTS(SELECT 1 FROM notifications n WHERE n.request_id=r.id AND n.type='new' AND n.processed_at IS NULL AND {recipient})"
        found=db.execute(f'SELECT r.*,{new} AS is_new FROM requests r WHERE {where} ORDER BY {sort} {direction},id DESC LIMIT 21 OFFSET %s',[*args,(page-1)*20]).fetchall()
        assigned={r[k] for r in found[:20] for k in ('fd_executor_id','uoso_executor_id') if r[k]} if g.user['role']!=1 else set()
        names={u['id']:u['full_name'] for u in db.execute('SELECT id,full_name FROM users WHERE id=ANY(%s)',(list(assigned),))}
    result=[record_payload(r,g.user['role']) for r in found[:20]]
    return jsonify(rows=result,hasNext=len(found)>20,names=names,page=page)

@app.get('/api/dashboard')
def dashboard():
    where,args=query_filters(g.user,{})
    with connection() as db:
        counts=db.execute(f'SELECT {STAGE} AS stage,count(*) AS n FROM requests r WHERE {where} GROUP BY 1',args).fetchall()
        latest=db.execute(f'SELECT r.* FROM requests r WHERE {where} ORDER BY registered_at DESC,id DESC LIMIT 6',args).fetchall()
        due=db.execute(f"SELECT count(*) AS n FROM requests r WHERE {where} AND return_date<=now() AT TIME ZONE 'Europe/Moscow'",args).fetchone()['n']
    stages={r['stage']:r['n'] for r in counts}
    return jsonify(counts=stages if g.user['role']!=1 else {},total=sum(stages.values()),due=due if g.user['role']!=1 else 0,rows=[record_payload(r,g.user['role']) for r in latest])

@app.get('/api/requests/<reqid>')
def read_request(reqid):
    with connection() as db:
        r=get_record(db,reqid)
        names={u['id']:u['full_name'] for u in db.execute('SELECT id,full_name FROM users WHERE id=ANY(%s)',([x for x in (r.get('fd_executor_id'),r.get('uoso_executor_id')) if x] if g.user['role']!=1 else [],))}
    return jsonify(row=record_payload(r,g.user['role']),names=names)

@app.post('/api/requests/<reqid>/read')
def mark_read(reqid):
    with connection() as db:
        r=get_record(db,reqid)
        db.execute("UPDATE notifications SET read_at=now() WHERE request_id=%s AND recipient_id=%s AND type='return'",(r['id'],g.user['id']))
    return jsonify(ok=True)

@app.patch('/api/requests/<reqid>')
def update_request(reqid):
    data=payload()
    with connection() as db:
        old=get_record(db,reqid,True)
        if data.get('revision')!=old['revision']:return jsonify(error='Заявка изменена другим пользователем. Откройте её заново.'),409
        changes=data.get('changes')
        if not isinstance(changes,dict):raise Invalid('Некорректные изменения')
        updated=apply_changes(old,changes,g.user,holidays(db))
        changed={k:v for k,v in updated.items() if v!=old.get(k) and k in COLUMNS}
        final=updated.get('decision_date') and updated.get('decision_date')!=old.get('decision_date')
        fd_final=updated.get('fd_status')=='Проведена' and old.get('fd_status')!='Проведена'
        if (final or fd_final) and not data.get('confirm'):return jsonify(error='Подтвердите завершение',confirmationRequired=True),409
        if changed:
            cols=list(changed)
            db.execute('UPDATE requests SET '+','.join(k+'=%s' for k in cols)+',revision=revision+1 WHERE id=%s',[*[changed[k] for k in cols],old['id']])
            audit(db,'Изменение заявки',old['id'],{k:old.get(k) for k in cols},changed)
            db.execute("UPDATE notifications SET processed_at=now(),read_at=now() WHERE request_id=%s AND type='new' AND processed_at IS NULL",(old['id'],))
            if updated.get('fd_status') and updated.get('fd_status')!=old.get('fd_status'):notify(db,old['id'],2,'Установлен Статус ФД: '+updated['fd_status'])
    return jsonify(ok=True)

@app.post('/api/requests/assign')
@roles(0,2,3)
def assign_requests():
    data=payload();idx=data.get('field');uid=data.get('userId');ids=data.get('ids')
    if idx not in (9,15) or not isinstance(ids,list) or not ids or len(ids)>9999:raise Invalid('Некорректное назначение')
    if g.user['role'] not in (0,3 if idx==9 else 2):raise PermissionError('Назначение недоступно')
    col=COLUMNS[idx];role=4 if idx==9 else 5
    with connection() as db:
        u=db.execute('SELECT * FROM users WHERE id=%s AND role=%s AND active',(uid,role)).fetchone()
        if not u:raise Invalid('Выберите активного исполнителя нужной роли')
        rows=db.execute('SELECT r.* FROM requests r WHERE request_id=ANY(%s) ORDER BY id FOR UPDATE',(ids,)).fetchall()
        if len(rows)!=len(set(ids)) or any(not can_view(r,g.user) or archive_due(r) for r in rows):raise PermissionError('Некоторые заявки недоступны')
        if idx==15 and any(not r['fd_status'] for r in rows):raise Invalid('Сначала установите Статус ФД')
        assigned=sum(r[col] is not None for r in rows)
        if assigned and not data.get('replace'):return jsonify(error='Уже назначен Исполнитель',assigned=assigned,confirmationRequired=True),409
        for r in rows:
            if r[col]==uid:continue
            db.execute(f'UPDATE requests SET {col}=%s,revision=revision+1 WHERE id=%s',(uid,r['id']))
            audit(db,'Назначение исполнителя',r['id'],{col:r[col]},{col:uid})
            notify(db,r['id'],role,'Вы назначены на заявку',uid)
    return jsonify(ok=True)

def import_rows(data,hs):
    rows=data.get('rows')
    if not isinstance(rows,list) or not rows:raise Invalid('Файл пустой')
    if len(rows)>9999:raise Invalid('Превышено максимальное количество строк: 9999')
    header=data.get('header')
    if not isinstance(header,list):raise Invalid('Структура файла не соответствует шаблону. Отсутствуют обязательные столбцы.')
    if header!=HEADERS[:8]:
        if len(header)<8:message='Отсутствуют обязательные столбцы.'
        elif sorted(map(str,header))==sorted(HEADERS[:8]):message='Нарушен порядок столбцов.'
        else:message='Неверные названия столбцов.'
        raise Invalid('Структура файла не соответствует шаблону. '+message)
    prepared=[]
    for values in rows:
        if not isinstance(values,list) or len(values)!=8:raise Invalid('Структура строки не соответствует шаблону')
        if all(v in ('',None) for v in values):raise Invalid('В файле обнаружена пустая строка. Загрузка остановлена.')
        if any(v in ('',None) for v in values):raise Invalid('Заполните обязательные поля')
        r=dict(zip(COLUMNS[:8],[str(v) for v in values]));r['registered_at']=parse_datetime(values[3]);calculate(r,hs);prepared.append(r)
    return prepared

@app.post('/api/imports/<action>')
@roles(0,1)
def imports(action):
    if action not in ('preview','commit'):raise Invalid('Неизвестная операция')
    data=payload()
    with connection() as db:
        rows=import_rows(data,holidays(db));ids=[r['request_id'] for r in rows]
        known={r['request_id'] for r in db.execute('SELECT request_id FROM requests WHERE request_id=ANY(%s)',(ids,))}
        created=0;skipped=0
        for r in rows:
            if r['request_id'] in known:skipped+=1;continue
            known.add(r['request_id'])
            if action=='commit':
                cols=COLUMNS[:9]+['calculated_decision_date','uploaded_by'];r['uploaded_by']=g.user['id']
                inserted=db.execute('INSERT INTO requests('+','.join(cols)+') VALUES ('+','.join(['%s']*len(cols))+') ON CONFLICT(request_id) DO NOTHING RETURNING id',[r.get(k) for k in cols]).fetchone()
                if not inserted:skipped+=1;continue
                audit(db,'Создана заявка',inserted['id'],None,{k:r[k] for k in COLUMNS[:8]})
                notify(db,inserted['id'],3,'Поступила новая заявка')
            created+=1
    return jsonify(created=created,skipped=skipped,errors=0)

@app.get('/api/export')
@roles(0,2,3)
def export():
    start,end=export_period(request.args.get('from'),request.args.get('to'))
    include=request.args.get('includeArchive')=='1'
    if include and g.user['role']!=0:raise PermissionError('Архив недоступен')
    where,args=query_filters(g.user,{'from':start.isoformat(),'to':end.isoformat()},include)
    wb=Workbook(write_only=True);ws=wb.create_sheet('Заявки');ws.append(HEADERS)
    with connection() as db:
        names={u['id']:u['full_name'] for u in db.execute('SELECT id,full_name FROM users')}
        with db.cursor(name='export_rows') as cur:
            cur.execute(f'SELECT r.* FROM requests r WHERE {where} ORDER BY registered_at,id',args)
            for r in cur:
                values=[]
                for i,col in enumerate(COLUMNS):
                    v=r.get(col)
                    if i in (9,15):v=names.get(v,'')
                    elif isinstance(v,datetime):v=v.strftime('%d.%m.%Y %H:%M:%S')
                    elif isinstance(v,date):v=v.strftime('%d.%m.%Y')
                    if isinstance(v,str) and v.startswith(('=','+','-','@')):
                        from openpyxl.cell import WriteOnlyCell
                        cell=WriteOnlyCell(ws,value=v);cell.data_type='s';v=cell
                    values.append(v)
                ws.append(values)
    out=io.BytesIO();wb.save(out);out.seek(0)
    return send_file(out,download_name=f'Заявки_с_{start:%d.%m.%Y}_по_{end:%d.%m.%Y}.xlsx',as_attachment=True,mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.post('/api/users')
@roles(0)
def create_user():
    d=payload();name=str(d.get('name','')).strip();login=str(d.get('login','')).strip();role=d.get('role')
    if not name or not login or type(role)!=int or role not in range(6):raise Invalid('Заполните ФИО, логин и роль')
    hashed=hash_password(d.get('password'))
    with connection() as db:
        u=db.execute('INSERT INTO users(full_name,login,role,password_hash) VALUES (%s,%s,%s,%s) RETURNING *',(name,login,role,hashed)).fetchone()
        audit(db,'Создание пользователя',after=public_user(u))
    return jsonify(user=public_user(u)),201

@app.patch('/api/users/<int:uid>')
@roles(0)
def update_user(uid):
    d=payload();name=str(d.get('name','')).strip();login=str(d.get('login','')).strip();role=d.get('role');active=d.get('active')
    if not name or not login or type(role)!=int or role not in range(6) or type(active)!=bool:raise Invalid('Некорректные данные пользователя')
    with connection() as db:
        old=db.execute('SELECT * FROM users WHERE id=%s FOR UPDATE',(uid,)).fetchone()
        if not old:raise Invalid('Пользователь не найден')
        if old['role']==0 and (role!=0 or not active):raise Invalid('Нельзя блокировать Администратора или менять его роль')
        if role!=old['role'] and db.execute('SELECT 1 FROM requests WHERE fd_executor_id=%s OR uoso_executor_id=%s LIMIT 1',(uid,uid)).fetchone():raise Invalid('Нельзя менять роль пользователя, связанного с заявками')
        u=db.execute('UPDATE users SET full_name=%s,login=%s,role=%s,active=%s WHERE id=%s RETURNING *',(name,login,role,active,uid)).fetchone()
        if not active or role!=old['role']:db.execute('DELETE FROM sessions WHERE user_id=%s',(uid,))
        audit(db,'Изменение пользователя',before=public_user(old),after=public_user(u))
    return jsonify(user=public_user(u))

@app.post('/api/users/<int:uid>/reset-password')
@roles(0)
def reset_password(uid):
    hashed=hash_password(payload().get('password'))
    with connection() as db:
        if not db.execute('UPDATE users SET password_hash=%s,must_change_password=true WHERE id=%s RETURNING id',(hashed,uid)).fetchone():raise Invalid('Пользователь не найден')
        db.execute('DELETE FROM sessions WHERE user_id=%s',(uid,))
        audit(db,'Сброс пароля',after={'userId':uid})
    return jsonify(ok=True)

@app.post('/api/holidays')
@roles(0)
def set_holidays():
    d=payload();values=d.get('dates',[])
    if not isinstance(values,list):raise Invalid('Ожидается список дат')
    dates={parse_date(v) for v in values}
    with connection() as db:
        # A single transaction makes calendar updates and dependent dates consistent.
        db.execute('SELECT pg_advisory_xact_lock(4857003)')
        if d.get('clear'):db.execute('DELETE FROM holidays')
        for value in dates:db.execute('INSERT INTO holidays(date) VALUES (%s) ON CONFLICT DO NOTHING',(value,))
        if d.get('remove'):db.execute('DELETE FROM holidays WHERE date=%s',(parse_date(d['remove']),))
        hs=holidays(db)
        with db.cursor(name='calendar_recalculate') as cur:
            cur.execute('SELECT r.* FROM requests r WHERE NOT '+ARCHIVED+' FOR UPDATE')
            for row in cur:
                calculate(row,hs)
                db.execute('UPDATE requests SET fd_due_date=%s,calculated_decision_date=%s,revision=revision+1 WHERE id=%s',(row['fd_due_date'],row['calculated_decision_date'],row['id']))
    return jsonify(ok=True)

@app.get('/api/templates')
@roles(0)
def templates():
    with connection() as db:rows=db.execute('SELECT t.id,t.filename,t.created_at,u.full_name FROM template_versions t JOIN users u ON u.id=t.uploaded_by ORDER BY id DESC LIMIT 100').fetchall()
    return jsonify(rows=rows)

@app.post('/api/templates')
@roles(0)
def upload_template():
    d=payload();filename=str(d.get('filename',''));suffix=Path(filename).suffix.lower()
    if suffix not in ('.xlsx','.csv'):raise Invalid('Недопустимый формат файла. Загрузите .xlsx или .csv')
    try:
        content=base64.b64decode(d.get('content',''),validate=True)
        if suffix=='.xlsx':
            wb=load_workbook(io.BytesIO(content),read_only=True,data_only=True);header=list(next(wb.active.values));wb.close()
        else:
            import csv
            raw=content.decode('utf-8-sig');dialect=csv.Sniffer().sniff(raw[:4096],delimiters=',;\t');header=next(csv.reader(io.StringIO(raw),dialect))
    except Exception:raise Invalid('Не удалось прочитать шаблон')
    if header!=HEADERS[:8]:raise Invalid('Шаблон должен содержать восемь эталонных столбцов в установленном порядке')
    with connection() as db:
        row=db.execute('INSERT INTO template_versions(filename,content,uploaded_by) VALUES (%s,%s,%s) RETURNING id',(Path(filename).name,content,g.user['id'])).fetchone()
        audit(db,'Загрузка шаблона',after={'id':row['id'],'filename':filename})
    return jsonify(ok=True)

@app.get('/api/templates/download')
@roles(0,1)
def download_template():
    with connection() as db:t=db.execute('SELECT filename,content FROM template_versions ORDER BY id DESC LIMIT 1').fetchone()
    if t:return send_file(io.BytesIO(bytes(t['content'])),download_name=t['filename'],as_attachment=True)
    wb=Workbook();wb.active.append(HEADERS[:8]);out=io.BytesIO();wb.save(out);out.seek(0)
    return send_file(out,download_name='Шаблон.xlsx',as_attachment=True)

@app.get('/api/audit')
@roles(0)
def action_history():
    try:page=max(1,int(request.args.get('page',1)))
    except ValueError:raise Invalid('Некорректная страница')
    with connection() as db:rows=db.execute("SELECT a.id,a.occurred_at,a.action,a.before_values,a.after_values,coalesce(u.full_name,'Система') AS full_name,r.request_id FROM action_log a LEFT JOIN users u ON u.id=a.user_id LEFT JOIN requests r ON r.id=a.request_id ORDER BY a.id DESC LIMIT 101 OFFSET %s",((page-1)*100,)).fetchall()
    return jsonify(rows=rows[:100],hasNext=len(rows)>100)

@app.get('/api/notifications')
def notification_list():
    try:page=max(1,int(request.args.get('page',1)))
    except ValueError:raise Invalid('Некорректная страница')
    where="n.processed_at IS NULL AND NOT "+ARCHIVED+" AND u.active AND (u.role IN (2,3) OR (u.role=4 AND r.fd_executor_id=u.id AND r.fd_status IS DISTINCT FROM 'Проведена') OR (u.role=5 AND r.uoso_executor_id=u.id AND r.decision_date IS NULL)) AND (n.type='new' OR (n.due_date=r.return_date AND n.due_date<=now() AT TIME ZONE 'Europe/Moscow'))"
    args=[]
    if g.user['role']!=0:where+=' AND n.recipient_id=%s';args.append(g.user['id'])
    base=' FROM notifications n JOIN requests r ON r.id=n.request_id JOIN users u ON u.id=n.recipient_id WHERE '+where
    with connection() as db:
        count=db.execute('SELECT count(*) AS n'+base+' AND n.read_at IS NULL',args).fetchone()['n']
        typ=request.args.get('type')
        if typ in ('new','return'):base+=' AND n.type=%s';args.append(typ)
        rows=db.execute('SELECT n.id,n.type,n.message,n.due_date,n.created_at,n.read_at,r.request_id,u.full_name AS recipient'+base+' ORDER BY n.id DESC LIMIT 21 OFFSET %s',[*args,(page-1)*20]).fetchall()
    for row in rows:
        for key in ('due_date','created_at','read_at'):
            if row[key]:row[key]=row[key].isoformat(timespec='seconds')
    return jsonify(rows=rows[:20],hasNext=len(rows)>20,unread=count)

@app.get('/')
def index():return send_from_directory(ROOT,'server.html')

@app.get('/<path:filename>')
def assets(filename):
    if filename not in ('server.js','server.css','xlsx.full.min.js'):return jsonify(error='Не найдено'),404
    return send_from_directory(ROOT,filename)
